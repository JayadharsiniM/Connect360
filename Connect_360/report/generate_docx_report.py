import os
import re
from bs4 import BeautifulSoup, NavigableString, Tag
import docx
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

REPORT_DIR = os.path.abspath(os.path.dirname(__file__))
PROJECT_ROOT = os.path.abspath(os.path.join(REPORT_DIR, ".."))
HTML_PATH = os.path.join(REPORT_DIR, "CONNECT360_PHASE1_PROJECT_REPORT.html")
DOCX_PATH = os.path.join(REPORT_DIR, "CONNECT360_PHASE1_PROJECT_REPORT.docx")

DIAGRAMS_DIR = os.path.join(REPORT_DIR, "diagrams_png")
EVAL_DIR = os.path.join(PROJECT_ROOT, "evaluation")

FIGURE_MAP = {
    "Figure 3.1": os.path.join(DIAGRAMS_DIR, "fig3_1_architecture.png"),
    "Figure 3.2": os.path.join(DIAGRAMS_DIR, "fig3_2_usecase.png"),
    "Figure 3.3": os.path.join(DIAGRAMS_DIR, "fig3_3_activity.png"),
    "Figure 3.4": os.path.join(DIAGRAMS_DIR, "fig3_4_dfd_context.png"),
    "Figure 3.5": os.path.join(DIAGRAMS_DIR, "fig3_5_dfd_modular.png"),
    "Figure 3.6": os.path.join(DIAGRAMS_DIR, "fig3_6_sequence.png"),
    "Figure 5.1": os.path.join(EVAL_DIR, "metric1_matching_latency", "metric1_matching_latency_final.png"),
    "Figure 5.2": os.path.join(EVAL_DIR, "metric2_rank_sensitivity", "metric2_rank_inversion.png"),
    "Figure 5.3": os.path.join(EVAL_DIR, "metric3_concurrency", "metric3_transactional_concurrency.png"),
    "Figure 5.4": os.path.join(EVAL_DIR, "metric4_spatial_routing", "metric4_road_vs_haversine_disparity.png"),
    "Figure 5.5": os.path.join(EVAL_DIR, "metric5_ai_safety", "metric5_confusion_matrix_alone.png"),
    "Figure 5.6": os.path.join(EVAL_DIR, "metric6_structured_output", "metric6_structured_output_performance.png"),
}

# Column widths dictionary by table type / caption substring
TABLE_WIDTHS = {
    "Comparative Evaluation": [Inches(1.2), Inches(1.2), Inches(1.2), Inches(1.2), Inches(1.22)],
    "Software Requirements": [Inches(1.5), Inches(2.1), Inches(2.42)],
    "Hardware Requirements": [Inches(1.5), Inches(1.8), Inches(2.72)],
    "DynamoDB Single-Table": [Inches(0.9), Inches(1.0), Inches(1.0), Inches(1.05), Inches(1.05), Inches(1.02)],
    "Multi-Criteria Weighting": [Inches(1.5), Inches(1.1), Inches(1.1), Inches(2.32)],
    "Matching Algorithm Execution Latency": [Inches(1.1), Inches(0.82), Inches(0.82), Inches(0.82), Inches(0.82), Inches(0.82), Inches(0.82)],
    "Rank Inversion & Sensitivity": [Inches(1.3), Inches(1.1), Inches(1.1), Inches(1.2), Inches(1.32)],
    "Transactional Concurrency": [Inches(0.7), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.67)],
    "Haversine vs. OSRM": [Inches(1.4), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.32)],
    "Defensive PII Redaction": [Inches(1.5), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7), Inches(1.02)],
    "Structured Output Schema": [Inches(1.4), Inches(0.75), Inches(0.75), Inches(0.75), Inches(0.75), Inches(0.75), Inches(0.87)],
    "ABBREVIATION": [Inches(1.6), Inches(4.42)],
    "TABLE OF CONTENTS": [Inches(1.3), Inches(3.92), Inches(0.8)],
    "LIST OF TABLES": [Inches(1.3), Inches(3.92), Inches(0.8)],
    "LIST OF FIGURES": [Inches(1.3), Inches(3.92), Inches(0.8)],
}

def format_cell_borders(cell, top=True, bottom=True, left=False, right=False, color="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f"""
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{'single' if top else 'none'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{'single' if left else 'none'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{'single' if bottom else 'none'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{'single' if right else 'none'}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tcBorders>
    """)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_footer_page_number(run):
    r = run._r
    fld1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fld2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fld3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    r.append(fld1)
    r.append(instr)
    r.append(fld2)
    r.append(fld3)

def add_runs_from_html(p, element, base_font_size=12, base_font_name="Times New Roman", italic_default=False, bold_default=False):
    for child in element.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if not text:
                continue
            r = p.add_run(text)
            r.font.name = base_font_name
            r.font.size = Pt(base_font_size)
            if italic_default:
                r.font.italic = True
            if bold_default:
                r.font.bold = True
        elif isinstance(child, Tag):
            tag_name = child.name.lower()
            child_text = child.get_text()
            if not child_text:
                continue
            r = p.add_run(child_text)
            r.font.name = base_font_name
            r.font.size = Pt(base_font_size)
            if tag_name in ["strong", "b"] or bold_default:
                r.font.bold = True
            if tag_name in ["em", "i"] or italic_default:
                r.font.italic = True
            if tag_name == "code":
                r.font.name = "Consolas"
                r.font.size = Pt(base_font_size - 0.5)

print("Starting DOCX generation...")

with open(HTML_PATH, "r", encoding="utf-8") as f:
    soup = BeautifulSoup(f.read(), "html.parser")

doc = Document()

# Set Normal Style
style_normal = doc.styles["Normal"]
style_normal.font.name = "Times New Roman"
style_normal.font.size = Pt(12)
style_normal.font.color.rgb = RGBColor(0, 0, 0)
style_normal.paragraph_format.line_spacing = 1.5
style_normal.paragraph_format.space_after = Pt(6)
style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# =============================================================================
# SECTION 1: PRELIMINARY PAGES (Title, Certificate, Abstract, Ack, TOC, Lists)
# =============================================================================
sec1 = doc.sections[0]
sec1.page_width = Mm(210)
sec1.page_height = Mm(297)
sec1.top_margin = Inches(1.0)
sec1.bottom_margin = Inches(1.0)
sec1.left_margin = Inches(1.25)
sec1.right_margin = Inches(1.0)
sec1.different_first_page_header_footer = True

sectPr1 = sec1._sectPr
pgNumType1 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="lowerRoman" w:start="1"/>')
sectPr1.append(pgNumType1)

footer1 = sec1.footer
p_f1 = footer1.paragraphs[0]
p_f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f1 = p_f1.add_run()
r_f1.font.name = "Times New Roman"
r_f1.font.size = Pt(11)
add_footer_page_number(r_f1)

# -----------------------------------------------------------------------------
# 1. TITLE PAGE
# -----------------------------------------------------------------------------
p_t1 = doc.add_paragraph()
p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_t1.paragraph_format.space_before = Pt(36)
p_t1.paragraph_format.space_after = Pt(18)
r = p_t1.add_run("CONNECT360 - AN INTELLIGENT AND SECURE PLATFORM FOR SKILLED LABOUR BOOKING")
r.font.name = "Times New Roman"
r.font.size = Pt(16)
r.font.bold = True

p_t2 = doc.add_paragraph()
p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_t2.paragraph_format.space_after = Pt(28)
r = p_t2.add_run("PHASE I REPORT")
r.font.name = "Times New Roman"
r.font.size = Pt(14)
r.font.bold = True

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(16)
r = p_sub.add_run("Submitted by")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

# Student Table
tbl_students = doc.add_table(rows=4, cols=2)
tbl_students.alignment = WD_TABLE_ALIGNMENT.CENTER
student_data = [
    ("JAYADHARSINI M", "230701127"),
    ("DHARSHINI R S", "230701076"),
    ("ENIYA B A", "230701085"),
    ("JAYAPRADHA P", "230701130"),
]
for idx, (name, reg) in enumerate(student_data):
    cell_l = tbl_students.cell(idx, 0)
    cell_r = tbl_students.cell(idx, 1)
    cell_l.width = Inches(2.2)
    cell_r.width = Inches(1.5)
    
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_l.paragraph_format.space_after = Pt(2)
    p_l.paragraph_format.line_spacing = 1.2
    r_l = p_l.add_run(name)
    r_l.font.name = "Times New Roman"
    r_l.font.size = Pt(12)
    r_l.font.bold = True

    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_r.paragraph_format.space_after = Pt(2)
    p_r.paragraph_format.line_spacing = 1.2
    r_r = p_r.add_run(reg)
    r_r.font.name = "Times New Roman"
    r_r.font.size = Pt(12)

# College Block
p_deg = doc.add_paragraph()
p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_deg.paragraph_format.space_before = Pt(36)
p_deg.paragraph_format.space_after = Pt(4)
r = p_deg.add_run("in partial fulfillment for the award of the degree of")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.italic = True

p_be = doc.add_paragraph()
p_be.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_be.paragraph_format.space_after = Pt(2)
r = p_be.add_run("BACHELOR OF ENGINEERING")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.font.bold = True

p_in = doc.add_paragraph()
p_in.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_in.paragraph_format.space_after = Pt(2)
r = p_in.add_run("IN")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.bold = True

p_cse = doc.add_paragraph()
p_cse.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cse.paragraph_format.space_after = Pt(40)
r = p_cse.add_run("COMPUTER SCIENCE AND ENGINEERING")
r.font.name = "Times New Roman"
r.font.size = Pt(13)
r.font.bold = True

p_col = doc.add_paragraph()
p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_col.paragraph_format.space_after = Pt(2)
r = p_col.add_run("RAJALAKSHMI ENGINEERING COLLEGE, CHENNAI\nANNA UNIVERSITY: CHENNAI 600 025\nMARCH 2026")
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r.font.bold = True

doc.add_page_break()

# -----------------------------------------------------------------------------
# 2. BONAFIDE CERTIFICATE
# -----------------------------------------------------------------------------
p_cert_h = doc.add_paragraph()
p_cert_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cert_h.paragraph_format.space_before = Pt(18)
p_cert_h.paragraph_format.space_after = Pt(4)
r = p_cert_h.add_run("ANNA UNIVERSITY: CHENNAI 600 025")
r.font.name = "Times New Roman"
r.font.size = Pt(14)
r.font.bold = True

p_cert_sub = doc.add_paragraph()
p_cert_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cert_sub.paragraph_format.space_after = Pt(24)
r = p_cert_sub.add_run("BONAFIDE CERTIFICATE")
r.font.name = "Times New Roman"
r.font.size = Pt(14)
r.font.bold = True

p_c1 = doc.add_paragraph()
p_c1.paragraph_format.line_spacing = 1.5
p_c1.paragraph_format.space_after = Pt(12)
r = p_c1.add_run('Certified that this project report titled ')
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r_bold = p_c1.add_run('"CONNECT360 - AN INTELLIGENT AND SECURE PLATFORM FOR SKILLED LABOUR BOOKING"')
r_bold.font.name = "Times New Roman"
r_bold.font.size = Pt(12)
r_bold.font.bold = True
r = p_c1.add_run(' is the bonafide work of ')
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r_team = p_c1.add_run('JAYADHARSINI M (230701127), DHARSHINI R S (230701076), ENIYA B A (230701085), and JAYAPRADHA P (230701130)')
r_team.font.name = "Times New Roman"
r_team.font.size = Pt(12)
r_team.font.bold = True
r = p_c1.add_run(' who carried out the project work under my supervision.')
r.font.name = "Times New Roman"
r.font.size = Pt(12)

p_c2 = doc.add_paragraph()
p_c2.paragraph_format.line_spacing = 1.5
p_c2.paragraph_format.space_after = Pt(36)
r = p_c2.add_run('Certified further that to the best of my knowledge the work reported herein does not form part of any other thesis or dissertation on the basis of which a degree or award was conferred on an earlier occasion on this or any other candidate. This project actively addresses ')
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r_sdg8 = p_c2.add_run('United Nations Sustainable Development Goal 8 (Decent Work and Economic Growth)')
r_sdg8.font.name = "Times New Roman"
r_sdg8.font.size = Pt(12)
r_sdg8.font.bold = True
r = p_c2.add_run(' and ')
r.font.name = "Times New Roman"
r.font.size = Pt(12)
r_sdg9 = p_c2.add_run('Goal 9 (Industry, Innovation, and Infrastructure)')
r_sdg9.font.name = "Times New Roman"
r_sdg9.font.size = Pt(12)
r_sdg9.font.bold = True
r = p_c2.add_run(' by establishing an equitable, transparent digital labor marketplace for independent technical service workers.')
r.font.name = "Times New Roman"
r.font.size = Pt(12)

# Signatures Table
tbl_sig = doc.add_table(rows=1, cols=2)
tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
c_l = tbl_sig.cell(0, 0)
c_r = tbl_sig.cell(0, 1)
c_l.width = Inches(3.0)
c_r.width = Inches(3.0)

p_sl = c_l.paragraphs[0]
p_sl.paragraph_format.line_spacing = 1.2
r = p_sl.add_run("SIGNATURE\n\n\n\n[TO BE PROVIDED]\nHEAD OF THE DEPARTMENT\n")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True
r_sub = p_sl.add_run("Professor,\nDepartment of Computer Science & Engg.,\nRajalakshmi Engineering College,\nThandalam, Chennai - 602 105.")
r_sub.font.name = "Times New Roman"
r_sub.font.size = Pt(10.5)

p_sr = c_r.paragraphs[0]
p_sr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_sr.paragraph_format.line_spacing = 1.2
r = p_sr.add_run("SIGNATURE\n\n\n\nMs. Divya M\nSUPERVISOR / PROJECT GUIDE\n")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True
r_sub = p_sr.add_run("Assistant Professor (CSE),\nDepartment of Computer Science & Engg.,\nRajalakshmi Engineering College,\nThandalam, Chennai - 602 105.")
r_sub.font.name = "Times New Roman"
r_sub.font.size = Pt(10.5)

p_viva = doc.add_paragraph()
p_viva.paragraph_format.space_before = Pt(30)
p_viva.paragraph_format.space_after = Pt(24)
r = p_viva.add_run("Submitted for the Phase-I Project Viva-Voce Examination held on ____________________ at Rajalakshmi Engineering College, Thandalam, Chennai.")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.italic = True

tbl_exam = doc.add_table(rows=1, cols=2)
tbl_exam.alignment = WD_TABLE_ALIGNMENT.CENTER
c_el = tbl_exam.cell(0, 0)
c_er = tbl_exam.cell(0, 1)
c_el.width = Inches(3.0)
c_er.width = Inches(3.0)

p_el = c_el.paragraphs[0]
r = p_el.add_run("INTERNAL EXAMINER")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

p_er = c_er.paragraphs[0]
p_er.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p_er.add_run("EXTERNAL EXAMINER")
r.font.name = "Times New Roman"
r.font.size = Pt(11)
r.font.bold = True

doc.add_page_break()

# -----------------------------------------------------------------------------
# 3. ABSTRACT
# -----------------------------------------------------------------------------
p_abs = doc.add_paragraph()
p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_abs.paragraph_format.space_before = Pt(12)
p_abs.paragraph_format.space_after = Pt(18)
r = p_abs.add_run("ABSTRACT")
r.font.name = "Times New Roman"
r.font.size = Pt(16)
r.font.bold = True

# Extract abstract paragraphs from HTML
abs_div = None
for div in soup.find_all("div", class_="chapter-header"):
    h2 = div.find("h2")
    if h2 and "ABSTRACT" in h2.get_text():
        abs_div = div
        break

if abs_div:
    curr = abs_div.next_sibling
    while curr:
        if isinstance(curr, Tag):
            if "page-break" in curr.get("class", []) or "chapter-header" in curr.get("class", []):
                break
            if curr.name == "p":
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(8)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs_from_html(p, curr, base_font_size=12)
        curr = curr.next_sibling

doc.add_page_break()

# -----------------------------------------------------------------------------
# 4. ACKNOWLEDGEMENT
# -----------------------------------------------------------------------------
p_ack = doc.add_paragraph()
p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ack.paragraph_format.space_before = Pt(12)
p_ack.paragraph_format.space_after = Pt(18)
r = p_ack.add_run("ACKNOWLEDGEMENT")
r.font.name = "Times New Roman"
r.font.size = Pt(16)
r.font.bold = True

ack_div = None
for div in soup.find_all("div", class_="chapter-header"):
    h2 = div.find("h2")
    if h2 and "ACKNOWLEDGEMENT" in h2.get_text():
        ack_div = div
        break

if ack_div:
    curr = ack_div.next_sibling
    while curr:
        if isinstance(curr, Tag):
            if "page-break" in curr.get("class", []) or "chapter-header" in curr.get("class", []):
                break
            if curr.name == "p":
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(8)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs_from_html(p, curr, base_font_size=12)
        curr = curr.next_sibling

doc.add_page_break()

# Helper to render an academic table from HTML table element
def render_academic_table(html_table, col_widths, caption_text=""):
    rows = html_table.find_all("tr")
    if not rows:
        return
    
    num_rows = len(rows)
    # determine max cols
    num_cols = max(len(r.find_all(["th", "td"])) for r in rows)
    
    tbl = doc.add_table(rows=num_rows, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    tblPr = tbl._tbl.tblPr
    borders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="333333"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="333333"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)

    for r_idx, r_elem in enumerate(rows):
        row = tbl.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        cells_html = r_elem.find_all(["th", "td"])
        is_header = bool(r_elem.find("th")) or (r_idx == 0)
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx, cell_html in enumerate(cells_html):
            if c_idx >= num_cols:
                break
            cell = row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            
            tcPr = cell._tc.get_or_add_tcPr()
            if is_header:
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>'))
            
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            # Alignment heuristics
            txt = cell_html.get_text().strip()
            cls = cell_html.get("class", [])
            style = cell_html.get("style", "")
            if "text-align: right" in style or "col-page" in cls:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif "text-align: center" in style or "col-num" in cls or is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            font_size = 10 if not is_header else 10.5
            add_runs_from_html(p, cell_html, base_font_size=font_size, bold_default=is_header)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)

# -----------------------------------------------------------------------------
# 5. TABLE OF CONTENTS, LIST OF TABLES, LIST OF FIGURES, LIST OF ABBREVIATIONS
# -----------------------------------------------------------------------------
prelim_sections = [
    ("TABLE OF CONTENTS", "TABLE OF CONTENTS"),
    ("LIST OF TABLES", "LIST OF TABLES"),
    ("LIST OF FIGURES", "LIST OF FIGURES"),
    ("LIST OF ABBREVIATIONS", "LIST OF ABBREVIATIONS"),
]

for title, key in prelim_sections:
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(18)
    r = p_h.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True

    sec_div = None
    for div in soup.find_all("div", class_="chapter-header"):
        h2 = div.find("h2")
        if h2 and key in h2.get_text():
            sec_div = div
            break

    if sec_div:
        curr = sec_div.next_sibling
        while curr:
            if isinstance(curr, Tag):
                if "page-break" in curr.get("class", []) or "chapter-header" in curr.get("class", []):
                    break
                if curr.name == "table":
                    widths = TABLE_WIDTHS.get(key, [Inches(1.5), Inches(3.72), Inches(0.8)])
                    render_academic_table(curr, widths, title)
            curr = curr.next_sibling

    doc.add_page_break()

# =============================================================================
# SECTION 2: BODY CHAPTERS (Arabic Numerals starting at 1)
# =============================================================================
sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
sec2.page_width = Mm(210)
sec2.page_height = Mm(297)
sec2.top_margin = Inches(1.0)
sec2.bottom_margin = Inches(1.0)
sec2.left_margin = Inches(1.25)
sec2.right_margin = Inches(1.0)
sec2.header.is_linked_to_previous = False
sec2.footer.is_linked_to_previous = False

sectPr2 = sec2._sectPr
pgNumType2 = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="decimal" w:start="1"/>')
sectPr2.append(pgNumType2)

footer2 = sec2.footer
p_f2 = footer2.paragraphs[0]
p_f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f2 = p_f2.add_run()
r_f2.font.name = "Times New Roman"
r_f2.font.size = Pt(11)
add_footer_page_number(r_f2)

# Collect all body chapters from HTML
chapter_nodes = []
current_ch = None

for elem in soup.body.children:
    if not isinstance(elem, Tag):
        continue
    classes = elem.get("class", [])
    if "chapter-header" in classes:
        h1 = elem.find("h1")
        h2 = elem.find("h2")
        # Check if it is Chapter 1..6 or REFERENCES
        title_text = elem.get_text().strip()
        if "CHAPTER" in title_text or "REFERENCES" in title_text:
            current_ch = {"elem": elem, "title": title_text, "children": []}
            chapter_nodes.append(current_ch)
    elif current_ch is not None:
        current_ch["children"].append(elem)

print(f"Collected {len(chapter_nodes)} body chapters to process.")

for ch_idx, ch in enumerate(chapter_nodes):
    print(f"Processing: {ch['title'].replace(chr(10), ' ')}")

    # Add Chapter Header
    header_div = ch["elem"]
    h1 = header_div.find("h1")
    h2 = header_div.find("h2")

    if h1:
        p1 = doc.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(18)
        p1.paragraph_format.space_after = Pt(4)
        p1.paragraph_format.keep_with_next = True
        r = p1.add_run(h1.get_text().strip().upper())
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True

    if h2:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(4)
        p2.paragraph_format.space_after = Pt(20)
        p2.paragraph_format.keep_with_next = True
        r = p2.add_run(h2.get_text().strip().upper())
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True

    # Process all elements inside chapter
    fig3_4_inserted = False

    for child in ch["children"]:
        if not isinstance(child, Tag):
            continue
        
        classes = child.get("class", [])

        # Heading 2 (e.g. 1.1 GENERAL)
        if child.name == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs_from_html(p, child, base_font_size=13.5, bold_default=True)

        # Heading 3 (e.g. 3.2.1 Software Requirements)
        elif child.name == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs_from_html(p, child, base_font_size=12, bold_default=True)

        # Heading 4
        elif child.name == "h4":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs_from_html(p, child, base_font_size=12, bold_default=True, italic_default=True)

        # Paragraph
        elif child.name == "p":
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs_from_html(p, child, base_font_size=12)

        # Equation
        elif "equation" in classes:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.2
            add_runs_from_html(p, child, base_font_size=11.5, italic_default=True)

        # Table caption
        elif "table-caption" in classes:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_runs_from_html(p, child, base_font_size=11, bold_default=True)

        # Academic Table
        elif child.name == "table" and "academic-table" in classes:
            cap = child.find_previous(class_="table-caption")
            cap_text = cap.get_text().strip() if cap else ""
            
            widths = None
            for k, w in TABLE_WIDTHS.items():
                if k.lower() in cap_text.lower():
                    widths = w
                    break
            if not widths:
                # default proportional
                num_cols = len(child.find_all("tr")[0].find_all(["th", "td"])) if child.find_all("tr") else 4
                w_per_col = Inches(6.02 / max(1, num_cols))
                widths = [w_per_col] * num_cols

            render_academic_table(child, widths, cap_text)

        # Figure Container
        elif "figure-container" in classes:
            caption_elem = child.find(class_="figure-caption")
            cap_text = caption_elem.get_text().strip() if caption_elem else ""

            # Check if Figure 3.5 is encountered, but Figure 3.4 was not yet inserted
            if "Figure 3.5" in cap_text and not fig3_4_inserted:
                fig3_4_path = FIGURE_MAP.get("Figure 3.4")
                if fig3_4_path and os.path.exists(fig3_4_path):
                    p_img34 = doc.add_paragraph()
                    p_img34.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img34.paragraph_format.space_before = Pt(12)
                    p_img34.paragraph_format.space_after = Pt(4)
                    r = p_img34.add_run()
                    r.add_picture(fig3_4_path, width=Inches(5.8))

                    p_cap34 = doc.add_paragraph()
                    p_cap34.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap34.paragraph_format.space_before = Pt(4)
                    p_cap34.paragraph_format.space_after = Pt(12)
                    r = p_cap34.add_run("Figure 3.4: Connect360 Data Flow Diagram (DFD Level 0 - Context Diagram)")
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
                    r.font.bold = True
                    fig3_4_inserted = True

            # Match Figure Image
            matched_fig = None
            for fig_key, fig_path in FIGURE_MAP.items():
                if fig_key in cap_text:
                    matched_fig = fig_path
                    break

            if matched_fig and os.path.exists(matched_fig):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                r = p_img.add_run()
                
                # Sizing: wide diagrams vs square confusion matrix
                if "confusion_matrix" in matched_fig:
                    r.add_picture(matched_fig, width=Inches(4.8))
                elif "fig3_" in matched_fig:
                    r.add_picture(matched_fig, width=Inches(5.8))
                else:
                    r.add_picture(matched_fig, width=Inches(5.8))

            if caption_elem:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(4)
                p_cap.paragraph_format.space_after = Pt(14)
                add_runs_from_html(p_cap, caption_elem, base_font_size=11, bold_default=True)

        # Bullet / Numbered Lists
        elif child.name in ["ul", "ol"]:
            for idx, li in enumerate(child.find_all("li")):
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.left_indent = Inches(0.35)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                
                if child.name == "ol":
                    prefix = f"{idx+1}.  "
                else:
                    prefix = "•  "
                
                r_pre = p.add_run(prefix)
                r_pre.font.name = "Times New Roman"
                r_pre.font.size = Pt(12)
                r_pre.font.bold = True
                
                add_runs_from_html(p, li, base_font_size=12)

        # Page Break
        elif "page-break" in classes:
            doc.add_page_break()

    # End of chapter page break (except for References)
    if ch_idx < len(chapter_nodes) - 1:
        doc.add_page_break()

# Save final document
print(f"Saving final report to {DOCX_PATH}...")
doc.save(DOCX_PATH)
print(f"Successfully generated DOCX report: {DOCX_PATH} ({os.path.getsize(DOCX_PATH)} bytes)")
